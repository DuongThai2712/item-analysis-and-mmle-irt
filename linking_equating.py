#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
linking_equating.py
===================
Linking (Stocking-Lord) + equating đề X -> đề Y trong thiết kế common-item
nonequivalent groups (Kolen & Brennan, mục 6.3.3), chạy riêng cho từng phần thi
TV / TA / TO / KH.

Cấu trúc file Excel đầu vào
---------------------------
  X, Y         : tham số câu hỏi (cột test, id, a, b; Y có thể đảo thứ tự cột -> đọc theo tên cột)
  Anchor       : cặp câu neo (cột X, Y = id câu ở đề X, đề Y)
  TV_X..KH_X   : bài làm đề X (1 đúng, 0 sai, -1 bỏ trống) + Raw/Theta/True
  TV_Y..KH_Y   : điểm đề Y (SBD, Raw, Theta, True)
  info_X/info_Y: SBD + Email người nộp bài (dùng email để tìm thí sinh làm cả hai đề -> kiểm tra link)

Quy ước (Kolen & Brennan, Eq. 6.2-6.5; Scale I = X, Scale J = Y)
-------------------------------------------------------------
  theta_Y = A * theta_X + B        a_Y = a_X / A        b_Y = A * b_X + B

Thang đo tham số
----------------
  * Tham số X được ước lượng bằng irt.py  -> P = expit(1.702 * a * (theta - b))
  * Tham số Y trong sheet Y               -> P = expit(1.0   * a * (theta - b))
    (đã kiểm tra: E[tcc(theta_Y)] ~ Raw_Y khi D = 1)
  Hai bên được quy về cùng thang logit (a_L = D * a) trước khi link; D chỉnh bằng --d-x / --d-y.
  Kết quả xuất ra có cả a thang D=1 (cùng thang đề Y) và a thang irt.py (= a_L / 1.702).

Stocking-Lord
-------------
  SLdiff(theta) = [ sum_j P_Y(A*theta+B ; a_Yj,b_Yj) - sum_j P_X(theta ; a_Xj,b_Xj) ]^2   (Eq. 6.13)
  SLcrit        = sum_theta w(theta) * SLdiff(theta)                                      (Eq. 6.15)
  Mặc định dùng dạng đối xứng của Kim & Kolen (2007): cộng thêm sai khác TCC tính trên lưới theta
  của nhóm Y (thang Y) với trọng số N(0,1); --one-sided để chỉ dùng nhóm X.
  Giá trị khởi đầu: mean/sigma và mean/mean (Eq. 6.8a/b, 6.9a).

Kết quả ghi trực tiếp vào workbook đầu vào
------------------------------------------
  * TV_X..KH_X : thêm cột Theta_X, Theta_new, True_new, Exclude, Pct_Raw, RawEq_Y
  * Link_Summary, Anchor_Link, X_Linked, Pct_TV.., Link_Check : sheet mới
  * plots/     : TCC (X trước / X sau link / Y), ICC + IIF của từng câu neo (legend có a, b)

Dùng:
    python linking_equating.py --input testequa.xlsx
"""
import argparse
import logging
import os
import shutil
import sys
import types
import warnings

import numpy as np
import pandas as pd
from scipy.optimize import minimize
from scipy.special import expit
from scipy.stats import norm

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

logging.getLogger("matplotlib.font_manager").setLevel(logging.ERROR)
warnings.filterwarnings("ignore", message=".*findfont.*")

HERE = os.path.dirname(os.path.abspath(__file__))
for _p in (HERE, os.getcwd()):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import irt  # noqa: E402  (irt.py của bạn: irt_probability, theta_estimate, true_score, D, MISSING)


def _import_item_plot():
    """
    Import item_plot (lazy, chỉ khi vẽ) để tránh import vòng: item_plot.py có `import linking_equating as le`.
    Đăng ký module hiện tại dưới tên 'linking_equating' để khi chạy `python linking_equating.py`
    item_plot dùng đúng module này thay vì nạp thêm một bản thứ hai.
    Nếu thiếu pygam/statsmodels thì dùng stub (các hàm vẽ ở đây không cần chúng).
    """
    sys.modules.setdefault("linking_equating", sys.modules[__name__])
    try:
        import item_plot as ip
        return ip
    except ImportError:
        for name in ("pygam", "statsmodels", "statsmodels.nonparametric",
                     "statsmodels.nonparametric.smoothers_lowess"):
            sys.modules.setdefault(name, types.ModuleType(name))
        sys.modules["pygam"].GAM = object
        sys.modules["pygam"].s = lambda *a, **k: None
        sys.modules["statsmodels.nonparametric.smoothers_lowess"].lowess = None
        import item_plot as ip
        return ip

# ----------------------------------------------------------------------------------------------
# Cấu hình chung
# ----------------------------------------------------------------------------------------------
TESTS = [  # (mã, tên, id đầu, id cuối)
    ("TV", "Tiếng Việt", 1, 30),
    ("TA", "Tiếng Anh", 31, 60),
    ("TO", "Toán", 61, 90),
    ("KH", "Tư duy khoa học", 91, 120),
]
SCALE_MAX = 300
COL_X = "#1f77b4"      # X gốc
COL_XL = "#d62728"     # X sau link
COL_Y = "#2ca02c"      # Y

HEADER_FILL = PatternFill("solid", start_color="DDEBF7")
HEADER_FONT = Font(bold=True)


def setup_style():
    try:
        import seaborn as sns
        sns.set_theme(style="whitegrid")
    except ImportError:
        pass
    plt.rcParams["font.family"] = "serif"
    plt.rcParams["font.serif"] = ["Times New Roman", "DejaVu Serif"]


# ----------------------------------------------------------------------------------------------
# Hàm IRT (thang logit: a_L = D * a)
# ----------------------------------------------------------------------------------------------
def p_logit(theta, a_L, b):
    theta = np.atleast_1d(np.asarray(theta, float))
    return expit(np.asarray(a_L, float)[None, :] * (theta[:, None] - np.asarray(b, float)[None, :]))


def tcc_logit(theta, a_L, b):
    return p_logit(theta, a_L, b).sum(axis=1)


def to_irt(a_L):
    """a thang logit (D=1) -> a của irt.py (D=1.702) để dùng irt.irt_probability / irt.cal_iif."""
    return np.asarray(a_L, float) / irt.D


# ----------------------------------------------------------------------------------------------
# Lưới theta + trọng số
# ----------------------------------------------------------------------------------------------
def make_grid(n=41, lim=4.0):
    return np.linspace(-lim, lim, n)


def make_weights(grid, kind="normal", theta_hat=None):
    if kind == "empirical" and theta_hat is not None and np.isfinite(theta_hat).sum() > 20:
        step = grid[1] - grid[0]
        edges = np.r_[grid - step / 2, grid[-1] + step / 2]
        w = np.histogram(theta_hat[np.isfinite(theta_hat)], bins=edges)[0].astype(float) + 1e-12
    else:
        w = norm.pdf(grid, 0.0, 1.0)
    return w / w.sum()


# ----------------------------------------------------------------------------------------------
# Mean/Sigma, Mean/Mean (Eq. 6.8-6.9) và Stocking-Lord (Eq. 6.13-6.15)
# ----------------------------------------------------------------------------------------------
def mean_sigma(ax, bx, ay, by):
    A = np.std(by, ddof=1) / np.std(bx, ddof=1)
    return A, np.mean(by) - A * np.mean(bx)


def mean_mean(ax, bx, ay, by):
    A = np.mean(ax) / np.mean(ay)
    return A, np.mean(by) - A * np.mean(bx)


def sl_components(A, B, ax, bx, ay, by, thx, wx, thy, wy, symmetric=True):
    """Trả về (SLcrit_X, SLcrit_Y). SLcrit_X: lưới nhóm X (thang X); SLcrit_Y: lưới nhóm Y (thang Y)."""
    d1 = tcc_logit(A * thx + B, ay, by) - tcc_logit(thx, ax, bx)
    c1 = float(np.sum(wx * d1 ** 2))
    c2 = 0.0
    if symmetric:
        d2 = tcc_logit(thy, ay, by) - tcc_logit((thy - B) / A, ax, bx)
        c2 = float(np.sum(wy * d2 ** 2))
    return c1, c2


def fit_stocking_lord(ax, bx, ay, by, thx, wx, thy, wy, symmetric=True, x0=None, fast=False):
    """
    Tối thiểu hoá SLcrit theo (A, B); tối ưu theo (log A, B) để A > 0.
    Mặc định đa khởi đầu (mean/sigma, mean/mean, (1,0)). Với bootstrap: x0=(log A, B) của nghiệm đầy đủ
    và fast=True (một khởi đầu ấm, dung sai nới) để nhanh hơn hàng chục lần.
    """
    def crit(p):
        A, B = np.exp(p[0]), p[1]
        c1, c2 = sl_components(A, B, ax, bx, ay, by, thx, wx, thy, wy, symmetric)
        return c1 + c2

    if x0 is not None:
        starts = [tuple(x0)]
    else:
        starts = []
        for f in (mean_sigma, mean_mean):
            A0, B0 = f(ax, bx, ay, by)
            if np.isfinite(A0) and A0 > 0:
                starts.append((np.log(A0), B0))
        starts.append((0.0, 0.0))
    opts = (dict(xatol=1e-6, fatol=1e-12, maxiter=500, maxfev=1000) if fast
            else dict(xatol=1e-9, fatol=1e-15, maxiter=4000, maxfev=8000))
    best = None
    for s in starts:
        r = minimize(crit, s, method="Nelder-Mead", options=opts)
        if best is None or r.fun < best.fun:
            best = r
    A, B = float(np.exp(best.x[0])), float(best.x[1])
    return A, B, float(best.fun)


def anchor_rmsd(A, B, ax, bx, ay, by, thy, wy):
    """RMSD (có trọng số) giữa ICC của từng câu neo X (sau link) và ICC câu Y, trên lưới theta thang Y."""
    px = p_logit(thy, ax / A, A * bx + B)
    py = p_logit(thy, ay, by)
    return np.sqrt(np.sum(wy[:, None] * (px - py) ** 2, axis=0))


# ----------------------------------------------------------------------------------------------
# Equipercentile (Kolen & Brennan, chương 2) - điểm nguyên 0..K
# ----------------------------------------------------------------------------------------------
def cdf_from_scores(scores, K):
    s = np.rint(np.asarray(scores, float)).astype(int)
    s = s[(s >= 0) & (s <= K)]
    f = np.bincount(s, minlength=K + 1)[:K + 1] / max(len(s), 1)
    return np.cumsum(f), f


def percentile_rank(x, F):
    """P(x) (đơn vị %): x có thể liên tục trong [-0.5, K+0.5]."""
    K = len(F) - 1
    x = np.atleast_1d(np.asarray(x, float))
    xs = np.clip(np.floor(x + 0.5).astype(int), 0, K)
    Fm1 = np.where(xs > 0, F[np.maximum(xs - 1, 0)], 0.0)
    pr = 100.0 * (Fm1 + (x - (xs - 0.5)) * (F[xs] - Fm1))
    return np.clip(pr, 0, 100)


def percentile_score(p, F):
    """Hàm ngược: điểm có percentile rank = p (%) theo phân phối F (công thức x_U của K&B)."""
    K = len(F) - 1
    p = np.atleast_1d(np.asarray(p, float))
    out = np.empty(len(p))
    for i, pi in enumerate(p):
        q = pi / 100.0
        if q <= 0:
            out[i] = -0.5
        elif q >= 1:
            out[i] = K + 0.5
        else:
            xs = int(np.searchsorted(F, q - 1e-12, side="left"))
            xs = min(xs, K)
            fm1 = F[xs - 1] if xs > 0 else 0.0
            den = F[xs] - fm1
            out[i] = (q - fm1) / den + xs - 0.5 if den > 0 else xs - 0.5
    return out


# ----------------------------------------------------------------------------------------------
# Đọc dữ liệu
# ----------------------------------------------------------------------------------------------
def _num(series):
    return pd.to_numeric(series.astype(str).str.replace(",", ".", regex=False), errors="coerce")


def read_workbook(path):
    sheets = pd.read_excel(path, sheet_name=None)
    for k in ("X", "Y", "Anchor"):
        if k not in sheets:
            raise SystemExit(f"Thiếu sheet '{k}' trong {path}")
    Xp, Yp = sheets["X"].copy(), sheets["Y"].copy()
    for d in (Xp, Yp):
        d["test"] = d["test"].ffill()
        d["id"] = d["id"].astype(int)
    anchors = sheets["Anchor"].dropna().astype(int)
    return sheets, Xp, Yp, anchors


# ----------------------------------------------------------------------------------------------
# Xử lý một phần thi
# ----------------------------------------------------------------------------------------------
def screen_anchors(ax_irt, bx, ay_L, by, a_min, b_max):
    """Loại câu neo có tham số không đáng tin: a của X chạm sàn, |b| quá lớn ở X hoặc Y."""
    keep, reason = [], []
    for a_x, b_x, a_y, b_y in zip(ax_irt, bx, ay_L, by):
        r = []
        if a_x <= a_min:
            r.append(f"a_X={a_x:.3f} chạm sàn")
        if abs(b_x) > b_max:
            r.append(f"|b_X|={abs(b_x):.2f}>{b_max:g}")
        if abs(b_y) > b_max:
            r.append(f"|b_Y|={abs(b_y):.2f}>{b_max:g}")
        keep.append(len(r) == 0)
        reason.append("; ".join(r))
    return np.array(keep), reason


def process_test(code, label, lo, hi, sheets, Xp, Yp, anchors, args):
    ids = list(range(lo, hi + 1))
    items = [f"Cau{i}" for i in ids]
    n_items = len(ids)

    xi = Xp[Xp.id.between(lo, hi)].set_index("id").loc[ids]
    yi = Yp[Yp.id.between(lo, hi)].set_index("id").loc[ids]
    aX_irt, bX = xi["a"].to_numpy(float), xi["b"].to_numpy(float)       # thang irt.py
    aX_L = args.d_x * aX_irt
    aY_L, bY = args.d_y * yi["a"].to_numpy(float), yi["b"].to_numpy(float)

    # ---------------- bài làm X: theta (MLE) + loại đúng hết / sai hết ----------------
    dX = sheets[f"{code}_X"].copy()
    cols = [f"Cau{i}" for i in ids]
    miss = [c for c in cols if c not in dX.columns]
    if miss:
        raise SystemExit(f"Sheet {code}_X thiếu cột {miss[:3]}...")
    U = dX[cols].to_numpy(float)
    Um = np.where(U == irt.MISSING, 0.0, U)                  # bỏ trống = 0 khi tính điểm thực / theta
    raw = (Um == 1).sum(axis=1)
    all_wrong, all_right = raw == 0, raw == n_items
    if args.use_sheet_theta:
        theta_x = dX["Theta"].to_numpy(float)
    else:
        theta_x = np.array(irt.theta_estimate(Um, a=aX_irt, b=bX, method="MLE"))
    at_bound = np.abs(theta_x) >= 6.0 - 1e-3
    exclude = np.where(all_wrong, "all_wrong",
               np.where(all_right, "all_correct",
               np.where(at_bound, "theta_bound", "")))
    ok = exclude == ""
    theta_x_ok = np.where(ok, theta_x, np.nan)
    theta_sheet_diff = float(np.nanmax(np.abs(np.where(ok, theta_x - dX["Theta"].to_numpy(float), np.nan))))

    # ---------------- điểm Y ----------------
    dY = sheets[f"{code}_Y"].copy()
    dY["Theta"] = _num(dY["Theta"])
    theta_y = dY["Theta"].to_numpy(float)

    # ---------------- neo ----------------
    A_ = anchors[(anchors.X >= lo) & (anchors.X <= hi)].reset_index(drop=True)
    bad = A_[~A_.Y.between(lo, hi)]
    if len(bad):
        print(f"  [CẢNH BÁO] {code}: neo có Y ngoài khối id {lo}-{hi}: {bad.to_dict('records')}")
    A_ = A_[A_.Y.between(lo, hi)].reset_index(drop=True)
    xa = Xp.set_index("id").loc[A_.X]
    ya = Yp.set_index("id").loc[A_.Y]
    ax_irt, bx_a = xa["a"].to_numpy(float), xa["b"].to_numpy(float)
    ax_a, ay_a, by_a = args.d_x * ax_irt, args.d_y * ya["a"].to_numpy(float), ya["b"].to_numpy(float)

    used, reason = screen_anchors(ax_irt, bx_a, ay_a, by_a, args.a_min, args.b_max)
    reason = list(reason)

    thx, thy = make_grid(args.quad_points, args.theta_lim), make_grid(args.quad_points, args.theta_lim)
    wx = make_weights(thx, args.weights, theta_x_ok)
    wy = make_weights(thy, args.weights, theta_y)
    sym = not args.one_sided

    def fit(mask):
        return fit_stocking_lord(ax_a[mask], bx_a[mask], ay_a[mask], by_a[mask], thx, wx, thy, wy, sym)

    if used.sum() < 3:
        raise SystemExit(f"{code}: chỉ còn {used.sum()} câu neo hợp lệ (<3), không link được.")
    A, B, crit = fit(used)

    # thanh lọc neo theo RMSD ICC (tuỳ chọn)
    if args.purify is not None:
        while used.sum() > args.min_anchor:
            r = anchor_rmsd(A, B, ax_a, bx_a, ay_a, by_a, thy, wy)
            r = np.where(used, r, -1)
            j = int(np.argmax(r))
            if r[j] <= args.purify:
                break
            used[j] = False
            reason[j] = f"purify: ICC RMSD={r[j]:.3f}>{args.purify:g}"
            A, B, crit = fit(used)

    c1, c2 = sl_components(A, B, ax_a[used], bx_a[used], ay_a[used], by_a[used], thx, wx, thy, wy, sym)
    A_ms, B_ms = mean_sigma(ax_a[used], bx_a[used], ay_a[used], by_a[used])
    A_mm, B_mm = mean_mean(ax_a[used], bx_a[used], ay_a[used], by_a[used])
    tcc_diff = tcc_logit(thy, ay_a[used], by_a[used]) - tcc_logit((thy - B) / A, ax_a[used], bx_a[used])
    tcc_rmsd = float(np.sqrt(np.sum(wy * tcc_diff ** 2)))
    tcc_maxd = float(np.max(np.abs(tcc_diff[np.abs(thy) <= 3])))
    rm_all = anchor_rmsd(A, B, ax_a, bx_a, ay_a, by_a, thy, wy)

    # bootstrap SE (lấy mẫu lại các câu neo đã dùng)
    se_A = se_B = np.nan
    if args.boot > 0:
        print(f"  bootstrap SE ({args.boot} lần)...", flush=True)
        rng = np.random.default_rng(args.seed)
        idx_used = np.flatnonzero(used)
        bs = []
        for _ in range(args.boot):
            s = rng.choice(idx_used, size=len(idx_used), replace=True)
            if np.std(bx_a[s]) < 1e-6:
                continue
            try:
                a_, b_, _ = fit_stocking_lord(ax_a[s], bx_a[s], ay_a[s], by_a[s], thx, wx, thy, wy, sym,
                                              x0=(np.log(A), B), fast=True)
                if np.isfinite(a_) and np.isfinite(b_):
                    bs.append((a_, b_))
            except Exception:
                continue
        if len(bs) > 10:
            bs = np.array(bs)
            se_A, se_B = float(bs[:, 0].std(ddof=1)), float(bs[:, 1].std(ddof=1))

    # ---------------- tham số X sau link ----------------
    aX_link, bX_link = aX_L / A, A * bX + B                       # thang Y (D=1)
    anchor_y_of = {int(x): int(y) for x, y in zip(A_.X, A_.Y)}
    used_x = {int(x) for x, u in zip(A_.X, used) if u}
    status, a_final, b_final, anc_y = [], aX_link.copy(), bX_link.copy(), []
    reason_of = {int(x): r for x, r in zip(A_.X, reason)}
    for k, i in enumerate(ids):
        if i in used_x:
            y_id = anchor_y_of[i]
            a_final[k], b_final[k] = args.d_y * Yp.set_index("id").loc[y_id, "a"], Yp.set_index("id").loc[y_id, "b"]
            status.append("Anchor (neo theo Y)")
            anc_y.append(y_id)
        elif i in anchor_y_of:
            status.append("Anchor bị loại: " + reason_of[i])
            anc_y.append(anchor_y_of[i])
        else:
            status.append("Link")
            anc_y.append(np.nan)
    x_linked = pd.DataFrame({
        "test": code, "id": ids, "Item": items, "Anchor_Y_id": anc_y, "Status": status,
        "a_X_sheet(D=1.702)": aX_irt, "b_X_sheet": bX,
        "a_X(D=1)": aX_L, "a_link(D=1)": aX_link, "b_link": bX_link,
        "a_final(D=1)": a_final, "b_final": b_final,
        "a_final(irt.py)": a_final / irt.D,
    })

    anchor_tbl = pd.DataFrame({
        "test": code, "X_id": A_.X, "Y_id": A_.Y, "Used": used, "Reason": reason,
        "a_X(sheet)": ax_irt, "b_X": bx_a, "a_X(D=1)": ax_a,
        "a_X_link(D=1)": ax_a / A, "b_X_link": A * bx_a + B,
        "a_Y(D=1)": ay_a, "b_Y": by_a,
        "d_a(link-Y)": ax_a / A - ay_a, "d_b(link-Y)": A * bx_a + B - by_a,
        "ICC_RMSD": rm_all,
    })

    # ---------------- theta mới + điểm thực mới (theo TCC đề Y) ----------------
    theta_new = np.where(ok, A * theta_x_ok + B, np.nan)
    y_params = pd.DataFrame({"a": aY_L / irt.D, "b": bY}, index=items)   # thang irt.py để dùng irt.true_score
    true_new = np.array([
        irt.true_score(theta_new[i], raw[i], pd.Series(Um[i], index=cols), y_params, scale_max=SCALE_MAX)
        for i in range(len(dX))
    ], dtype=float)
    # câu hỏi bỏ ngỏ: người theta_bound không có theta tin cậy -> true_score trả NaN (giữ NaN)

    # ---------------- equipercentile ----------------
    Kraw = n_items
    keep_pct = ~(all_wrong | all_right) if args.pct_drop_extremes else np.ones(len(dX), bool)
    F_x_raw, f_x_raw = cdf_from_scores(dX["Raw"].to_numpy(float)[keep_pct], Kraw)
    F_y_raw, f_y_raw = cdf_from_scores(dY["Raw"].to_numpy(float), Kraw)
    xs = np.arange(0, Kraw + 1)
    pr_x = percentile_rank(xs, F_x_raw)
    pr_y = percentile_rank(xs, F_y_raw)
    x2y = percentile_score(pr_x, F_y_raw)
    y2x = percentile_score(pr_y, F_x_raw)
    pct_p = np.arange(1, 100)
    F_x_true, _ = cdf_from_scores(dX["True"].to_numpy(float)[keep_pct], SCALE_MAX)
    F_y_true, _ = cdf_from_scores(dY["True"].to_numpy(float), SCALE_MAX)
    tn = true_new[np.isfinite(true_new) & keep_pct]
    F_xn_true, _ = cdf_from_scores(tn, SCALE_MAX)
    pct_tbl = pd.DataFrame({
        "Percentile": pct_p,
        "Raw_X": percentile_score(pct_p, F_x_raw), "Raw_Y": percentile_score(pct_p, F_y_raw),
        "True_X_old": percentile_score(pct_p, F_x_true),
        "True_X_new": percentile_score(pct_p, F_xn_true),
        "True_Y": percentile_score(pct_p, F_y_true),
    })
    raw_tbl = pd.DataFrame({
        "Raw": xs, "n_X": np.rint(f_x_raw * keep_pct.sum()).astype(int), "PR_X(%)": pr_x, "Raw_Y_equiv": x2y,
        "n_Y": np.rint(f_y_raw * len(dY)).astype(int), "PR_Y(%)": pr_y, "Raw_X_equiv": y2x,
    })
    person_cols = {
        "Theta_X": theta_x_ok, "Theta_new": theta_new, "True_new": true_new, "Exclude": exclude,
        "Pct_Raw": percentile_rank(raw.astype(float), F_x_raw),
        "RawEq_Y": percentile_score(percentile_rank(raw.astype(float), F_x_raw), F_y_raw),
    }

    summary = {
        "Test": code, "Tên": label, "N_anchor": len(A_), "N_used": int(used.sum()),
        "A": A, "B": B, "SE_A(boot)": se_A, "SE_B(boot)": se_B,
        "SLcrit": crit, "SLcrit_X": c1, "SLcrit_Y": c2,
        "TCC_RMSD(anchor)": tcc_rmsd, "TCC_maxdiff(|θ|<=3)": tcc_maxd,
        "A_MeanSigma": A_ms, "B_MeanSigma": B_ms, "A_MeanMean": A_mm, "B_MeanMean": B_mm,
        "mean_b_X(anchor)": float(np.mean(bx_a[used])), "mean_b_Y(anchor)": float(np.mean(by_a[used])),
        "sd_b_X(anchor)": float(np.std(bx_a[used], ddof=1)), "sd_b_Y(anchor)": float(np.std(by_a[used], ddof=1)),
        "mean_a_X(D=1)": float(np.mean(ax_a[used])), "mean_a_Y(D=1)": float(np.mean(ay_a[used])),
        "N_person": len(dX), "N_all_wrong": int(all_wrong.sum()), "N_all_correct": int(all_right.sum()),
        "N_theta_bound": int((exclude == "theta_bound").sum()), "N_linked": int(ok.sum()),
        "theta_X_mean": float(np.nanmean(theta_x_ok)), "theta_X_sd": float(np.nanstd(theta_x_ok, ddof=1)),
        "theta_new_mean": float(np.nanmean(theta_new)), "theta_new_sd": float(np.nanstd(theta_new, ddof=1)),
        "theta_Y_mean": float(np.nanmean(theta_y)), "theta_Y_sd": float(np.nanstd(theta_y, ddof=1)),
        "True_X_old_mean": float(dX["True"].mean()), "True_new_mean": float(np.nanmean(true_new)),
        "True_Y_mean": float(dY["True"].mean()),
        "max|Theta_sheet - Theta_X|": theta_sheet_diff,
    }

    return dict(code=code, label=label, ids=ids, items=items, A=A, B=B, summary=summary,
                x_linked=x_linked, anchor_tbl=anchor_tbl, anchors=A_, used=used,
                aX_irt=aX_irt, bX=bX, aX_L=aX_L, aY_L=aY_L, bY=bY,
                aX_link=aX_link, bX_link=bX_link, ax_a=ax_a, bx_a=bx_a, ay_a=ay_a, by_a=by_a,
                ax_irt=ax_irt, raw_tbl=raw_tbl, pct_tbl=pct_tbl, person_cols=person_cols,
                theta_x_ok=theta_x_ok, theta_new=theta_new, true_new=true_new, raw=raw,
                dX=dX, dY=dY, crit=crit, tcc_rmsd=tcc_rmsd)


# ----------------------------------------------------------------------------------------------
# Kiểm tra bằng thí sinh làm cả hai đề (ghép bằng email)
# ----------------------------------------------------------------------------------------------
def link_check(sheets, results):
    if "info_X" not in sheets or "info_Y" not in sheets:
        return None, None
    ix, iy = sheets["info_X"].copy(), sheets["info_Y"].copy()
    for d in (ix, iy):
        d["SBD"] = pd.to_numeric(d["SBD"], errors="coerce")
        d["key"] = d["Email"].astype(str).str.strip().str.lower()
    m = ix.merge(iy, on="key", suffixes=("_X", "_Y"))
    rows = []
    for r in results:
        dX, dY = r["dX"].set_index("SBD"), r["dY"].set_index("SBD")
        thx = pd.Series(r["theta_x_ok"], index=r["dX"]["SBD"].to_numpy())
        thn = pd.Series(r["theta_new"], index=r["dX"]["SBD"].to_numpy())
        tn = pd.Series(r["true_new"], index=r["dX"]["SBD"].to_numpy())
        for _, p in m.iterrows():
            sx, sy = p["SBD_X"], p["SBD_Y"]
            if sx in dX.index and sy in dY.index:
                rows.append({
                    "test": r["code"], "Email": p["key"], "SBD_X": int(sx), "SBD_Y": int(sy),
                    "Raw_X": dX.loc[sx, "Raw"], "Raw_Y": dY.loc[sy, "Raw"],
                    "Theta_X": thx.loc[sx], "Theta_new": thn.loc[sx], "Theta_Y": _num(pd.Series([dY.loc[sy, "Theta"]])).iloc[0],
                    "True_X_old": dX.loc[sx, "True"], "True_new": tn.loc[sx], "True_Y": dY.loc[sy, "True"],
                })
    chk = pd.DataFrame(rows)
    if chk.empty:
        return chk, None
    out = []
    for code, g in chk.groupby("test", sort=False):
        g = g.dropna(subset=["Theta_X", "Theta_new", "Theta_Y"])
        if len(g) < 3:
            continue
        out.append({
            "test": code, "n": len(g),
            "corr(Theta_X,Theta_Y)": np.corrcoef(g.Theta_X, g.Theta_Y)[0, 1],
            "corr(Theta_new,Theta_Y)": np.corrcoef(g.Theta_new, g.Theta_Y)[0, 1],
            "mean(Theta_X-Theta_Y)": (g.Theta_X - g.Theta_Y).mean(),
            "mean(Theta_new-Theta_Y)": (g.Theta_new - g.Theta_Y).mean(),
            "RMSE(Theta_X,Theta_Y)": np.sqrt(((g.Theta_X - g.Theta_Y) ** 2).mean()),
            "RMSE(Theta_new,Theta_Y)": np.sqrt(((g.Theta_new - g.Theta_Y) ** 2).mean()),
            "mean(True_X_old-True_Y)": (g.True_X_old - g.True_Y).mean(),
            "mean(True_new-True_Y)": (g.True_new - g.True_Y).mean(),
        })
    return chk, pd.DataFrame(out)


# ----------------------------------------------------------------------------------------------
# Ghi vào workbook
# ----------------------------------------------------------------------------------------------
def _clean(v):
    if isinstance(v, (np.floating, float)):
        return None if not np.isfinite(v) else float(v)
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.bool_,)):
        return bool(v)
    return v


def write_df(ws, df, r0=1, c0=1, num_fmt="0.0000", title=None, note=None):
    """Ghi DataFrame vào ws từ ô (r0, c0); trả về dòng kế tiếp còn trống."""
    r = r0
    if title:
        ws.cell(r, c0, title).font = Font(bold=True, size=12)
        r += 1
    if note:
        ws.cell(r, c0, note).alignment = Alignment(wrap_text=False)
        r += 1
    for j, name in enumerate(df.columns):
        c = ws.cell(r, c0 + j, str(name))
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            c = ws.cell(r + i, c0 + j, _clean(v))
            if isinstance(c.value, float):
                c.number_format = num_fmt
    for j, name in enumerate(df.columns):
        col = get_column_letter(c0 + j)
        w = max([len(str(name))] + [len(f"{x:.4f}") if isinstance(x, float) else len(str(x)) for x in df.iloc[:, j].head(200)])
        ws.column_dimensions[col].width = min(max(ws.column_dimensions[col].width or 0, w + 2, 9), 60)
    return r + len(df) + 2


def fresh_sheet(wb, name):
    if name in wb.sheetnames:
        del wb[name]
    return wb.create_sheet(name)


def write_person_cols(ws, cols):
    """Thêm (hoặc ghi đè nếu đã có) các cột kết quả vào cuối sheet bài làm đề X."""
    head = {str(c.value): c.column for c in ws[1] if c.value is not None}
    last = max(head.values())
    ref = ws.cell(1, last)
    for name, values in cols.items():
        if name in head:
            c = head[name]
        else:
            last += 1
            c = last
            hc = ws.cell(1, c, name)
            hc.font, hc.fill, hc.alignment = HEADER_FONT, HEADER_FILL, Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(c)].width = 12
        for i, v in enumerate(values, start=2):
            cell = ws.cell(i, c, _clean(v) if not isinstance(v, str) else (v or None))
            if isinstance(cell.value, float):
                cell.number_format = "0.0000" if name in ("Theta_X", "Theta_new") else "0.00"


def write_workbook(path_in, path_out, results, chk, chk_sum, args):
    wb = openpyxl.load_workbook(path_in)

    # ---- cột mới trong các sheet bài làm đề X ----
    for r in results:
        write_person_cols(wb[f"{r['code']}_X"], r["person_cols"])

    # ---- Link_Summary ----
    ws = fresh_sheet(wb, "Link_Summary")
    summ = pd.DataFrame([r["summary"] for r in results]).set_index("Test").T.reset_index()
    summ.columns = ["Chỉ số"] + [r["code"] for r in results]
    notes = [
        "Stocking-Lord (Kolen & Brennan, 6.3.3): tìm A, B tối thiểu hoá SLcrit = Σ w(θ)[τ_Y(Aθ+B) − τ_X(θ)]² (+ phần đối xứng trên lưới nhóm Y nếu không dùng --one-sided).",
        "Công thức chuyển thang: θ_Y = A·θ_X + B ; a_Y = a_X / A ; b_Y = A·b_X + B (Scale I = X, Scale J = Y).",
        f"Thang a: X (irt.py) D={args.d_x:g} ; Y D={args.d_y:g}. Hai bên quy về logit (a·D) trước khi link; 'a(D=1)' = thang đề Y.",
        f"Trọng số theta: {args.weights}, lưới {args.quad_points} điểm trong [-{args.theta_lim:g}, {args.theta_lim:g}]. Chế độ: {'một phía (nhóm X)' if args.one_sided else 'đối xứng (Kim & Kolen 2007)'}.",
        f"Câu neo bị loại trước khi link nếu a_X <= {args.a_min:g} (chạm sàn ước lượng) hoặc |b| > {args.b_max:g} ở X hoặc Y"
        + (f"; thanh lọc thêm theo ICC RMSD > {args.purify:g}" if args.purify is not None else "; không thanh lọc theo RMSD") + ".",
        f"SE(A), SE(B): bootstrap {args.boot} lần trên các câu neo đã dùng (seed={args.seed}).",
        "Theta_new = A·Theta_X + B; Theta_X được ước lượng lại (MLE, bỏ trống = 0) bằng tham số X trong sheet 'X'. Loại khỏi quy đổi: all_wrong, all_correct, theta_bound (|θ|≥6).",
        "True_new = TCC của đề Y tại Theta_new / số câu × 300 (irt.true_score). Người sai hết → 0, đúng hết → 300.",
    ]
    nr = write_df(ws, summ, 1, 1, title="Kết quả linking X → Y theo từng phần thi", num_fmt="0.0000")
    ws.column_dimensions["A"].width = 30
    ws.cell(nr, 1, "Ghi chú").font = Font(bold=True)
    for k, t in enumerate(notes, start=1):
        ws.cell(nr + k, 1, t)
    ws.freeze_panes = "B3"

    # ---- Anchor_Link ----
    ws = fresh_sheet(wb, "Anchor_Link")
    anc = pd.concat([r["anchor_tbl"] for r in results], ignore_index=True)
    write_df(ws, anc, 1, 1, title="Chi tiết câu neo: tham số X sau link so với Y",
             note="Used=False: câu neo bị loại khỏi tính A, B (xem Reason). ICC_RMSD: chênh lệch ICC có trọng số N(0,1) giữa X sau link và Y.")
    ws.freeze_panes = "A4"

    # ---- X_Linked ----
    ws = fresh_sheet(wb, "X_Linked")
    xl = pd.concat([r["x_linked"] for r in results], ignore_index=True)
    write_df(ws, xl, 1, 1, title="Tham số mới của đề X sau link (thang Y); câu neo hợp lệ được neo theo tham số Y",
             note="a(D=1), b: thang đề Y. a_final(irt.py) = a_final(D=1)/1.702 để dùng trực tiếp với irt.py. Câu neo hợp lệ: a_final, b_final = tham số Y.")
    ws.freeze_panes = "A4"

    # ---- Pct_<test> ----
    for r in results:
        ws = fresh_sheet(wb, f"Pct_{r['code']}")
        nr = write_df(ws, r["raw_tbl"], 1, 1, num_fmt="0.00",
                      title=f"{r['label']} – quy đổi equipercentile điểm thô X ↔ Y (Kolen & Brennan, chương 2)",
                      note="PR = percentile rank (%), Raw_Y_equiv = điểm thô đề Y có cùng bách phân vị với điểm đề X (và ngược lại). Giả định hai nhóm tương đương.")
        write_df(ws, r["pct_tbl"], nr, 1, num_fmt="0.00",
                 title="Điểm tại từng bách phân vị 1–99 của X và Y (True_X_new: điểm thực quy về thang Y sau link)")

    # ---- Link_Check ----
    if chk is not None and len(chk):
        ws = fresh_sheet(wb, "Link_Check")
        write_df(ws, chk, 1, 1, num_fmt="0.0000",
                 title="Thí sinh làm cả hai đề (ghép theo email): so Theta/True sau link với Theta/True của đề Y",
                 note="Chỉ để kiểm tra độ hợp lý của link (mẫu nhỏ, hai đề khác nhau nên có sai số đo + hiệu ứng luyện tập).")
        if chk_sum is not None:
            write_df(ws, chk_sum, 1, chk.shape[1] + 2, num_fmt="0.0000", title="Tổng hợp theo phần thi")

    if os.path.abspath(path_in) == os.path.abspath(path_out):
        bak = os.path.splitext(path_in)[0] + ".orig.xlsx"
        if args.backup and not os.path.exists(bak):
            shutil.copy2(path_in, bak)
            print(f"Đã sao lưu file gốc -> {bak}")
    wb.save(path_out)
    print(f"Đã ghi kết quả vào {path_out}")

# ----------------------------------------------------------------------------------------------
# main
# ----------------------------------------------------------------------------------------------
def parse_args(argv=None):
    ap = argparse.ArgumentParser(description="Linking Stocking-Lord + equating đề X -> đề Y")
    ap.add_argument("--input", default="testequa.xlsx", help="file Excel đầu vào")
    ap.add_argument("--output", default=None, help="file Excel đầu ra (mặc định: ghi thẳng vào file đầu vào)")
    ap.add_argument("--plots-dir", default=None, help="thư mục lưu hình (mặc định: <thư mục file>/plots)")
    ap.add_argument("--no-plots", action="store_true")
    ap.add_argument("--no-backup", dest="backup", action="store_false", help="không tạo <tên>.orig.xlsx khi ghi đè")
    ap.add_argument("--d-x", type=float, default=irt.D, help="hằng số D của tham số X (mặc định 1.702 như irt.py)")
    ap.add_argument("--d-y", type=float, default=1.0, help="hằng số D của tham số Y (mặc định 1.0)")
    ap.add_argument("--weights", choices=["normal", "empirical"], default="normal",
                    help="trọng số theta: N(0,1) (như MMLE) hoặc phân phối theta ước lượng")
    ap.add_argument("--quad-points", type=int, default=41)
    ap.add_argument("--theta-lim", type=float, default=4.0)
    ap.add_argument("--one-sided", action="store_true", help="SL một phía (chỉ lưới nhóm X), mặc định đối xứng")
    ap.add_argument("--a-min", type=float, default=0.06, help="loại neo có a_X (thang irt.py) <= ngưỡng này")
    ap.add_argument("--b-max", type=float, default=4.0, help="loại neo có |b| > ngưỡng này ở X hoặc Y")
    ap.add_argument("--purify", type=float, default=None,
                    help="nếu đặt: lần lượt bỏ neo có ICC RMSD > ngưỡng (vd 0.15) rồi link lại")
    ap.add_argument("--min-anchor", type=int, default=8, help="số neo tối thiểu khi --purify")
    ap.add_argument("--boot", type=int, default=200, help="số lần bootstrap SE(A), SE(B); 0 = tắt")
    ap.add_argument("--seed", type=int, default=2024)
    ap.add_argument("--pct-keep-extremes", dest="pct_drop_extremes", action="store_false",
                    help="giữ thí sinh X sai hết / đúng hết khi tính bách phân vị "
                         "(mặc định bỏ, vì đề Y không có điểm 0 và theo yêu cầu loại các trường hợp này)")
    ap.add_argument("--use-sheet-theta", action="store_true",
                    help="dùng cột Theta có sẵn trong sheet thay vì ước lượng lại từ bài làm")
    return ap.parse_args(argv)

def style_axes_zero(ax):
    """Trục theo mẫu drawTTC: trục x, y đi qua 0; bỏ khung phải và trên."""
    ax.spines["left"].set_position("zero")
    ax.spines["bottom"].set_position("zero")
    ax.spines["right"].set_color("none")
    ax.spines["top"].set_color("none")
    ax.xaxis.set_ticks_position("bottom")
    ax.yaxis.set_ticks_position("left")
    ax.yaxis.set_label_coords(-0.04, 0.5)   # nhãn trục y ở mép trái, không đè lên đường cong
 

def main(argv=None):
    args = parse_args(argv)
    path_in = args.input
    path_out = args.output or path_in
    plots_dir = args.plots_dir or os.path.join(os.path.dirname(os.path.abspath(path_out)), "plots")
    setup_style()

    sheets, Xp, Yp, anchors = read_workbook(path_in)
    print(f"Đọc {path_in}: {len(Xp)} câu X, {len(Yp)} câu Y, {len(anchors)} câu neo")

    results = []
    for code, label, lo, hi in TESTS:
        print(f"\n== {code} ({label}) ==")
        r = process_test(code, label, lo, hi, sheets, Xp, Yp, anchors, args)
        s = r["summary"]
        print(f"  neo dùng {s['N_used']}/{s['N_anchor']} | A = {r['A']:.4f} (SE {s['SE_A(boot)']:.3f}), "
              f"B = {r['B']:.4f} (SE {s['SE_B(boot)']:.3f}) | SLcrit = {r['crit']:.5f} | TCC RMSD = {s['TCC_RMSD(anchor)']:.3f}")
        print(f"  mean/sigma: A = {s['A_MeanSigma']:.4f}, B = {s['B_MeanSigma']:.4f} | "
              f"mean/mean: A = {s['A_MeanMean']:.4f}, B = {s['B_MeanMean']:.4f}")
        print(f"  loại khỏi quy đổi theta: sai hết {s['N_all_wrong']}, đúng hết {s['N_all_correct']}, "
              f"theta chạm cận {s['N_theta_bound']}; còn {s['N_linked']}/{s['N_person']}")
        print(f"  True TB: X cũ {s['True_X_old_mean']:.1f} -> X mới {s['True_new_mean']:.1f} | Y {s['True_Y_mean']:.1f}")
        results.append(r)

    chk, chk_sum = link_check(sheets, results)
    if chk_sum is not None:
        print("\nKiểm tra bằng thí sinh làm cả hai đề:")
        print(chk_sum.round(3).to_string(index=False))

    write_workbook(path_in, path_out, results, chk, chk_sum, args)

    if not args.no_plots:
        ip = _import_item_plot()
        os.makedirs(plots_dir, exist_ok=True)
        for r in results:
            ip.plot_tcc_anchor(r, plots_dir)
            ip.plot_anchor_curves(r, plots_dir)
        fig, axes = plt.subplots(ncols=2, nrows=2, figsize=(24, 12))
        for ax, r in zip(axes.flat, results):
            ip.plot_tcc_anchor(r, plots_dir, ax=ax)
        fig.suptitle("Đường cong đặc tính bài thi (TCC)", fontsize=14)
        fig.tight_layout()
        fig.savefig(os.path.join(plots_dir, "TCC_all.png"), dpi=150, bbox_inches="tight")
        plt.close(fig)
        print(f"Đã lưu toàn bộ hình vào {plots_dir}")


if __name__ == "__main__":
    main()