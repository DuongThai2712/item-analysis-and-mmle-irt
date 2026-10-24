# -*- coding: utf-8 -*-
"""
EXAM MIXER v4
=============
So với v3 (bản bạn gửi) - giữ nguyên toàn bộ logic trộn/tokenize, thêm:

1. Tách thành 3 bước: load_project -> make_exam_plan -> render_exam, để có thể
   đổi đáp án đúng gốc rồi render lại mà KHÔNG phải trộn lại (dùng cho web server).
2. Hỗ trợ 0 / 1 / nhiều đáp án đúng trên 1 câu (trước đây chỉ giữ đáp án đầu tiên).
3. Xuất file mapping (CSV, phân cách ';'):
     mỗi dòng 1 mã đề, cột CauN = câu gốc N, ô "{STT câu ở đề sao}-{4 chữ cái}",
     vd "2-DCA*B": 4 chữ cái = đáp án GỐC ứng với vị trí A,B,C,D của đề sao;
     dấu * ngay sau chữ cái = đáp án đúng (có thể nhiều *);
     không có * -> tất cả thí sinh được điểm câu đó.
4. Xuất dapan.xlsx: sheet "Theo câu gốc" (Gốc | từng mã đề; STT A B C D; ô đúng tô
   vàng) + các sheet chữ cái như file đáp án cũ.
"""
import re
import os
import sys
import time
import random

class BraceMismatchError(Exception):
    """Ngoặc {} không cân trong file nội dung -> báo lỗi rõ ràng thay vì
    âm thầm sinh ra output rác (cắt chuỗi tại vị trí -1)."""
    pass


def find_matching_brace(s, start, context_label=None):
    depth = 0
    i = start
    n = len(s)
    while i < n:
        c = s[i]
        if c == '\\':
            # \{  \}  \%  là ký hiệu hiển thị ký tự thật (control symbol),
            # KHÔNG phải dấu mở/đóng nhóm hay bắt đầu comment của TeX -> bỏ
            # qua, không tính vào độ sâu và không kích hoạt bỏ-qua-comment.
            # (\% dùng để in ra dấu % thật, ví dụ "120\%." -- nếu không loại
            # trừ, dấu % này sẽ bị hiểu nhầm là bắt đầu comment và nuốt mất
            # phần còn lại của dòng, bao gồm cả dấu } cần đóng.)
            # \\  (2 dấu \\ liên tiếp, lệnh xuống dòng) phải được nhận diện
            # là 1 token riêng trước, để không hiểu nhầm dấu \\ thứ 2 đứng
            # ngay trước 1 dấu {/% thật là escape của chính nó.
            if i + 1 < n and s[i + 1] == '\\':
                i += 2
                continue
            if i + 1 < n and s[i + 1] in '{}%':
                i += 2
                continue
            i += 1
            continue
        # Bỏ qua comment LaTeX (từ '%' chưa bị escape đến hết dòng) khi đếm
        # ngoặc, vì LaTeX cũng bỏ qua toàn bộ nội dung đó -- nếu không, một
        # dấu {/} lẻ nằm trong ghi chú sẽ bị đếm nhầm và báo lỗi giả.
        if c == '%':
            nl = s.find('\n', i)
            i = nl if nl != -1 else n
            continue
        if c == '{':
            depth += 1
        elif c == '}':
            depth -= 1
            if depth == 0:
                return i
        i += 1
    snippet_start = max(0, start - 40)
    snippet = s[snippet_start:start + 80].replace('\n', '\\n')
    where = f" trong {context_label}" if context_label else ""
    raise BraceMismatchError(
        f"Không tìm thấy dấu '}}' khớp với '{{' tại vị trí {start}{where}.\n"
        f"  -> Kiểm tra lại số ngoặc {{}} quanh đoạn: ...{snippet}...\n"
        f"  (Có thể do thiếu dấu }} khi gõ/copy câu hỏi, hoặc dùng \\{{ \\}} không cân trong công thức toán.)"
    )


def find_outside_comment(content, needle, start=0):
    """Giống content.find(needle, start) nhưng bỏ qua mọi vị trí nằm trong
    LaTeX comment (từ '%' chưa escape đến hết dòng) và bỏ qua các escape
    \\{ \\} \\% \\\\ giống hệt logic trong find_matching_brace.

    Đây là bản sửa cốt lõi cho lỗi "câu bị comment vẫn tính là câu thật":
    trước đây tokenize_content dùng content.find() thô nên 1 dòng dạng
    "%\\errorq{...}" (dùng làm mẫu, đã bị comment) vẫn bị nhận diện là lệnh
    \\errorq thật, khiến số câu bị đếm dư (vd 121 thay vì 120).
    """
    i = start
    n = len(content)
    L = len(needle)
    while i < n:
        # QUAN TRỌNG: phải kiểm tra khớp needle TRƯỚC khi xử lý escape/comment
        # bên dưới. needle luôn bắt đầu bằng '\' (vd \group, \errorq{...}), nếu
        # kiểm tra khớp SAU nhánh xử lý '\\' thì vị trí bắt đầu bằng '\' sẽ luôn
        # bị "nhảy qua" trước khi kịp so khớp -> không bao giờ tìm thấy gì cả.
        if content[i:i + L] == needle:
            return i
        c = content[i]
        if c == '\\':
            # \{ \} \% \\ là 2 ký tự escape/ký hiệu, nhảy qua cả cặp để dấu %
            # trong \% không bị hiểu nhầm là bắt đầu comment ở vòng lặp sau.
            if i + 1 < n and content[i + 1] in '\\{}%':
                i += 2
                continue
            i += 1
            continue
        if c == '%':
            nl = content.find('\n', i)
            i = nl if nl != -1 else n
            continue
        i += 1
    return -1


def parse_choice_in_text(text):
    m = re.search(r'\\choice(?:One|Two|Four)', text)
    if not m:
        return None
    macro = m.group()
    pos = m.end()
    args = []
    for _ in range(4):
        while pos < len(text) and text[pos] in ' \t\n\r':
            pos += 1
        if pos >= len(text) or text[pos] != '{':
            return None
        end = find_matching_brace(text, pos)
        arg_content = text[pos + 1:end]
        after = end + 1
        has_percent = False
        tmp = after
        # SỬA: trước đây chỉ bỏ qua ' \t' (cùng dòng) nên nếu dấu % bị lệch
        # sang dòng khác so với dấu '}' thì không nhận ra được đáp án đúng
        # (=> answer_table ra "None" dù người dùng đã đánh dấu % trong source).
        # Nay bỏ qua cả xuống dòng '\n' '\r' -- an toàn vì nếu ký tự non-space
        # tiếp theo không phải '%' thì `after` vẫn giữ nguyên giá trị cũ
        # (end + 1), không ảnh hưởng tới việc parse các tham số sau.
        while tmp < len(text) and text[tmp] in ' \t\r\n':
            tmp += 1
        if tmp < len(text) and text[tmp] == '%':
            has_percent = True
            after = tmp + 1
        args.append((arg_content, has_percent))
        pos = after
    pre = text[:m.start()]
    post = text[pos:]
    return pre, macro, args, post


def shuffle_choices(args, rng):
    letters = ['A', 'B', 'C', 'D']
    correct_idx = next((i for i, (_, hp) in enumerate(args) if hp), None)
    indices = list(range(4))
    rng.shuffle(indices)
    new_args, new_letter = [], None
    for new_pos, old_idx in enumerate(indices):
        content, _ = args[old_idx]
        is_correct = (old_idx == correct_idx)
        if is_correct:
            new_letter = letters[new_pos]
        new_args.append((content, is_correct))
    return new_args, new_letter


def get_correct_letter(args):
    letters = ['A', 'B', 'C', 'D']
    for i, (_, hp) in enumerate(args):
        if hp:
            return letters[i]
    return None


def count_marked_correct(args):
    """Đếm số đáp án đang được đánh dấu đúng (%) trong 1 câu -- dùng để cảnh
    báo nếu = 0 (thiếu dấu %) hoặc > 1 (đánh dấu nhầm nhiều đáp án)."""
    return sum(1 for _, hp in args if hp)


def rebuild_choice(macro, args):
    parts = [macro]
    for content, hp in args:
        parts.append('{' + content + '}%' if hp else '{' + content + '}')
    return '\n'.join(parts)


def apply_shuffle_to_q_text(q_text, rng, do_shuffle=True):
    result = parse_choice_in_text(q_text)
    if result is None:
        return q_text, None
    pre, macro, args, post = result
    if do_shuffle:
        new_args, letter = shuffle_choices(args, rng)
    else:
        new_args, letter = args, get_correct_letter(args)
    return pre + rebuild_choice(macro, new_args) + post, letter


def get_errorq_correct_letter(text):
    m = re.search(r'%([A-D])', text)
    return m.group(1) if m else None


_STOP_CMDS = [r'\reading', r'\engread', r'\q{', r'\q ', r'\errorq{',
              r'\randomgroup', r'\group', r'\noshuffle']


def _is_bare_cmd(content, i, name):
    L = len(name)
    if content[i:i + L] != name:
        return False
    nxt = i + L
    return not (nxt < len(content) and content[nxt].isalpha())


def tokenize_content(content):
    segments = []
    i, n = 0, len(content)

    while i < n:
        if content[i:i + 8] == r'\reading':
            j = i + 8
            while j < n and content[j] in ' \t\n':
                j += 1
            if j < n and content[j] == '{':
                end1 = find_matching_brace(content, j)
                j2 = end1 + 1
                while j2 < n and content[j2] in ' \t\n':
                    j2 += 1
                if j2 < n and content[j2] == '{':
                    end2 = find_matching_brace(content, j2)
                    seg_text = content[i:end2 + 1]
                    second_arg = content[j2:end2 + 1]
                    if r'\q' in second_arg:
                        segments.extend(parse_embedded_reading(seg_text, second_arg))
                    else:
                        segments.append({'type': 'reading_header', 'text': seg_text})
                    i = end2 + 1
                    continue

        if content[i:i + 8] == r'\engread':
            j = i + 8
            while j < n and content[j] in ' \t\n':
                j += 1
            if j < n and content[j] == '{':
                end1 = find_matching_brace(content, j)
                j2 = end1 + 1
                while j2 < n and content[j2] in ' \t\n':
                    j2 += 1
                if j2 < n and content[j2] == '{':
                    end2 = find_matching_brace(content, j2)
                    seg_text = content[i:end2 + 1]
                    second_arg = content[j2:end2 + 1]
                    # Giống \reading: nếu bên trong \engread{...}{...} có \q{...} thật sự
                    # (một số bài đọc TA nhúng thẳng \q/\choiceFour vào #2 thay vì chỉ
                    # dùng \item), thì phải tách ra thành các câu riêng để đánh số/xáo
                    # đáp án/vào bảng đáp án -- nếu không sẽ mất trắng các câu này.
                    if r'\q{' in second_arg or r'\q ' in second_arg:
                        segments.extend(parse_embedded_engread(seg_text, second_arg))
                    else:
                        segments.append({'type': 'engread_header', 'text': seg_text})
                    i = end2 + 1
                    continue

        if _is_bare_cmd(content, i, r'\randomgroup'):
            segments.append({'type': 'randomgroup_marker', 'text': r'\randomgroup'})
            i += len(r'\randomgroup')
            continue

        if _is_bare_cmd(content, i, r'\group'):
            segments.append({'type': 'group_marker', 'text': r'\group'})
            i += len(r'\group')
            continue

        if _is_bare_cmd(content, i, r'\noshuffle'):
            segments.append({'type': 'noshuffle_marker', 'text': r'\noshuffle'})
            i += len(r'\noshuffle')
            continue

        if content[i:i + 2] == r'\q' and (i + 2 >= n or content[i + 2] in ' \t\n{'):
            j = i + 2
            while j < n and content[j] in ' \t\n':
                j += 1
            if j < n and content[j] == '{':
                end = find_matching_brace(content, j)
                q_text = content[i:end + 1]
                if not re.search(r'\\choice(?:One|Two|Four)', q_text):
                    next_q_pos = n
                    for nc in [r'\q{', r'\q ', r'\errorq{']:
                        # SỬA: dùng find_outside_comment thay vì content.find
                        # thô, để không bị 1 ví dụ mẫu đã comment ra phía sau
                        # "kéo" phạm vi câu hỏi hiện tại đi sai.
                        nidx = find_outside_comment(content, nc, end + 1)
                        if nidx != -1 and nidx < next_q_pos:
                            next_q_pos = nidx
                    between = content[end + 1:next_q_pos]
                    cm = re.search(r'\\choice(?:One|Two|Four)', between)
                    if cm:
                        cm_pos = end + 1 + cm.start() + len(cm.group())
                        for _ in range(4):
                            while cm_pos < n and content[cm_pos] in ' \t\n':
                                cm_pos += 1
                            if cm_pos < n and content[cm_pos] == '{':
                                ce = find_matching_brace(content, cm_pos)
                                tmp = ce + 1
                                while tmp < n and content[tmp] in ' \t':
                                    tmp += 1
                                cm_pos = tmp + 1 if (tmp < n and content[tmp] == '%') else ce + 1
                        q_text = content[i:cm_pos]
                        end = cm_pos - 1
                segments.append({'type': 'q', 'text': q_text})
                i = end + 1
                continue

        if content[i:i + 8] == r'\errorq{':
            pos = i + 7
            full = r'\errorq'
            for _ in range(9):
                while pos < n and content[pos] in ' \t\n\r':
                    pos += 1
                if pos < n and content[pos] == '{':
                    end = find_matching_brace(content, pos)
                    full += content[pos:end + 1]
                    pos = end + 1
                else:
                    break
            eol = content.find('\n', pos)
            if eol == -1:
                eol = n
            full += content[pos:eol]
            segments.append({'type': 'errorq', 'text': full})
            i = eol
            continue

        next_pos = n
        for cmd in _STOP_CMDS:
            # SỬA CHÍNH: dùng find_outside_comment thay vì content.find thô.
            # Đây là chỗ trực tiếp gây ra lỗi "121 câu thay vì 120" -- 1 dòng
            # \errorq{...} bị comment ra (%\errorq{...}) trước đây vẫn được
            # find() thô "tìm thấy" và bị coi là điểm bắt đầu 1 lệnh thật,
            # khiến nó bị tokenize thành 1 câu errorq "ma".
            idx = find_outside_comment(content, cmd, i + 1)
            if idx != -1 and idx < next_pos:
                next_pos = idx
        segments.append({'type': 'other', 'text': content[i:next_pos]})
        i = next_pos

    return segments


def parse_embedded_reading(full_seg, second_arg):
    inner = second_arg[1:-1]
    q_pos = inner.find(r'\q{')
    if q_pos == -1:
        q_pos = inner.find(r'\q ')
    if q_pos == -1:
        return [{'type': 'reading_header', 'text': full_seg}]
    preamble = inner[:q_pos]
    num_match = re.search(r'\\reading\{(\d+)\}', full_seg)
    num = num_match.group(1) if num_match else '?'
    header_text = r'\reading{' + num + r'}{' + preamble + r'}'
    return [{'type': 'reading_header', 'text': header_text}] + tokenize_content(inner[q_pos:])


def parse_embedded_engread(full_seg, second_arg):
    """Tương tự parse_embedded_reading nhưng cho \\engread{n1}{n2}: nếu #2 (n2)
    chứa \\q{...} thật (không chỉ \\item thường), tách phần mở đầu (đề bài/đoạn
    văn + \\item không có \\q) làm header tĩnh, phần còn lại tokenize như file
    nội dung bình thường để mỗi câu được đếm số & xáo đáp án riêng."""
    inner = second_arg[1:-1]
    q_pos = inner.find(r'\q{')
    if q_pos == -1:
        q_pos = inner.find(r'\q ')
    if q_pos == -1:
        return [{'type': 'engread_header', 'text': full_seg}]
    preamble = inner[:q_pos]
    num_match = re.search(r'\\engread\{(\d+)\}', full_seg)
    num = num_match.group(1) if num_match else '?'
    header_text = r'\engread{' + num + r'}{' + preamble + r'}'
    return [{'type': 'engread_header', 'text': header_text}] + tokenize_content(inner[q_pos:])


def build_units(segments):
    units = []
    prefix_buf = []
    group_open, group_id, cur_group = False, 0, None
    rgroup_open, rgroup_id, cur_rgroup = False, 0, None
    seen_q = False
    file_level_noshuffle = False
    pending_noshuffle = False

    def flush_prefix():
        nonlocal prefix_buf
        if prefix_buf:
            units.append({'kind': 'fixed_text', 'prefix': prefix_buf, 'seg': None,
                          'group_tag': None, 'rgroup_tag': cur_rgroup, 'noshuffle': False})
            prefix_buf = []

    for seg in segments:
        t = seg['type']
        if t == 'randomgroup_marker':
            rgroup_open = not rgroup_open
            if rgroup_open:
                rgroup_id += 1
                cur_rgroup = rgroup_id
            else:
                cur_rgroup = None
            continue
        if t == 'group_marker':
            group_open = not group_open
            if group_open:
                group_id += 1
                cur_group = group_id
            else:
                cur_group = None
            continue
        if t == 'noshuffle_marker':
            if not seen_q:
                file_level_noshuffle = True
            elif prefix_buf and not ''.join(s['text'] for s in prefix_buf).strip():
                for prev in reversed(units):
                    if prev['kind'] in ('q', 'errorq'):
                        prev['noshuffle'] = True
                        break
                    if prev['kind'] not in ('fixed_text',):
                        break
            else:
                pending_noshuffle = True
            continue
        if t in ('q', 'errorq'):
            seen_q = True
            units.append({'kind': t, 'prefix': prefix_buf, 'seg': seg,
                          'group_tag': cur_group, 'rgroup_tag': cur_rgroup,
                          'noshuffle': pending_noshuffle})
            pending_noshuffle = False
            prefix_buf = []
        elif t in ('reading_header', 'engread_header'):
            flush_prefix()
            units.append({'kind': 'header', 'prefix': [], 'seg': seg,
                          'group_tag': None, 'rgroup_tag': cur_rgroup, 'noshuffle': False})
        else:
            prefix_buf.append(seg)

    flush_prefix()
    return units, file_level_noshuffle


def shuffle_within_groups(units, rng):
    i, n = 0, len(units)
    while i < n:
        gid = units[i]['group_tag']
        if gid is None:
            i += 1
            continue
        j = i
        while j < n and units[j]['group_tag'] == gid:
            j += 1
        sub = units[i:j]
        rng.shuffle(sub)
        units[i:j] = sub
        i = j


def make_chunks(units):
    chunks, i, n = [], 0, len(units)
    while i < n:
        gid = units[i]['group_tag']
        if gid is not None:
            j = i
            while j < n and units[j]['group_tag'] == gid:
                j += 1
            chunks.append(units[i:j])
            i = j
        else:
            chunks.append([units[i]])
            i += 1
    return chunks


def shuffle_chunks_by_randomgroup(chunks, rng):
    out = list(chunks)
    i, n = 0, len(out)
    while i < n:
        rg = out[i][0]['rgroup_tag']
        if rg is None:
            i += 1
            continue
        j = i
        while j < n and out[j][0]['rgroup_tag'] == rg:
            j += 1
        sub = out[i:j]
        rng.shuffle(sub)
        out[i:j] = sub
        i = j
    return [u for c in out for u in c]


def parse_structure_file(text):
    segments, i, n = [], 0, len(text)
    while i < n:
        if _is_bare_cmd(text, i, r'\randomgroup'):
            segments.append(('randomgroup', None))
            i += len(r'\randomgroup')
            continue
        if text[i:i + 11] == r'\inputspace':
            j = i + 11
            while j < n and text[j] in ' \t':
                j += 1
            if j < n and text[j] == '{':
                end = find_matching_brace(text, j)
                segments.append(('inputspace', text[j + 1:end]))
                i = end + 1
                continue
        next_pos = n
        for cmd in [r'\randomgroup', r'\inputspace']:
            idx = find_outside_comment(text, cmd, i + 1)
            if idx != -1 and idx < next_pos:
                next_pos = idx
        segments.append(('text', text[i:next_pos]))
        i = next_pos
    return segments


def build_structure_chunks(segments):
    chunks, prefix = [], []
    rgroup_open, rgroup_id, cur_rg = False, 0, None
    for kind, val in segments:
        if kind == 'randomgroup':
            rgroup_open = not rgroup_open
            if rgroup_open:
                rgroup_id += 1
                cur_rg = rgroup_id
            else:
                cur_rg = None
            continue
        if kind == 'text':
            prefix.append(val)
        else:
            chunks.append({'prefix': ''.join(prefix), 'path': val, 'rgroup_tag': cur_rg})
            prefix = []
    if prefix:
        chunks.append({'prefix': ''.join(prefix), 'path': None, 'rgroup_tag': None})
    return chunks


def shuffle_structure_chunks(chunks, rng):
    out = list(chunks)
    i, n = 0, len(out)
    while i < n:
        rg = out[i]['rgroup_tag']
        if rg is None:
            i += 1
            continue
        j = i
        while j < n and out[j]['rgroup_tag'] == rg:
            j += 1
        sub = out[i:j]
        rng.shuffle(sub)
        out[i:j] = sub
        i = j
    return out


def rel_path(path):
    return path + '.tex'


def regions_by_tag(seq, tag_of):
    regions, i, n = [], 0, len(seq)
    while i < n:
        tag = tag_of(seq[i])
        if tag is None:
            i += 1
            continue
        j = i
        while j < n and tag_of(seq[j]) == tag:
            j += 1
        regions.append((tag, i, j))
        i = j
    return regions


def _label_unit(u):
    if u['kind'] in ('q', 'errorq'):
        return f"câu#{u.get('orig_q_num', '?')}"
    if u['kind'] == 'header':
        return '(tiêu đề bài đọc)'
    return '(text)'


def _label_chunk(chunk):
    if len(chunk) == 1:
        return _label_unit(chunk[0])
    labels = [_label_unit(u) for u in chunk if u['kind'] in ('q', 'errorq')]
    return 'nhóm[' + ','.join(labels) + ']'


def snapshot_regions(seq, tag_of, label_of):
    out = {}
    order = []
    for tag, s, e in regions_by_tag(seq, tag_of):
        if tag not in out:
            order.append(tag)
        out.setdefault(tag, []).extend(label_of(x) for x in seq[s:e])
    return out, order


def _short_snippet(text, maxlen=90):
    flat = re.sub(r'\s+', ' ', text).strip()
    return (flat[:maxlen] + '…') if len(flat) > maxlen else flat


class ProgressTracker:
    def __init__(self, exam_code, exam_index, total_exams):
        self.exam_code = exam_code
        self.exam_index = exam_index
        self.total_exams = total_exams
        self.lines = []
        self.warnings = []
        self.n_struct_regions = 0
        self.n_group_regions = 0
        self.n_rgroup_regions = 0
        self.n_files = 0
        self.t0 = time.time()

    def log(self, text):
        self.lines.append(text)

    def warn(self, text):
        """Cảnh báo mức độ 'cần người dùng kiểm tra lại file gốc' -- in ra
        console ngay (để không bị bỏ sót) và ghi vào log."""
        self.warnings.append(text)
        print(f"    ⚠ {text}")
        self.log(f"⚠ CẢNH BÁO: {text}")

    def struct_regions(self, before_map, order):
        self.n_struct_regions += len(order)
        for tag in order:
            self.log(f"[Cấu trúc] Cụm #{tag}: {' -> '.join(before_map[tag])}")

    def struct_regions_after(self, after_map, order):
        for tag in order:
            self.log(f"           đảo thành: {' -> '.join(after_map[tag])}")

    def file_start(self, i, total, path):
        self.n_files += 1
        print(f"  [{self.exam_index}/{self.total_exams}] ({i}/{total}) đang trộn {path} ...")

    def group_region(self, path, tag, before, after):
        self.n_group_regions += 1
        self.log(f"[{path}] \\group #{tag}: {' | '.join(before)}")
        self.log(f"{' ' * (len(path) + 2)}   -> {' | '.join(after)}")

    def rgroup_region(self, path, tag, before, after):
        self.n_rgroup_regions += 1
        self.log(f"[{path}] \\randomgroup #{tag}: {' | '.join(before)}")
        self.log(f"{' ' * (len(path) + 2)}      -> {' | '.join(after)}")

    def finish(self, n_questions, out_dir):
        elapsed = time.time() - self.t0
        print(f"  Mã đề {self.exam_code} (STT {self.exam_index}): {n_questions} câu | "
              f"{self.n_struct_regions} cụm cấu trúc, {self.n_group_regions} \\group, "
              f"{self.n_rgroup_regions} \\randomgroup đã đảo, "
              f"{len(self.warnings)} cảnh báo -> {out_dir}  ({elapsed:.2f}s)")

    def text(self):
        out = [f"NHẬT KÝ ĐẢO ĐỀ — mã đề {self.exam_code} (STT {self.exam_index})",
               "=" * 60]
        if self.warnings:
            out.append(f"\n*** CÓ {len(self.warnings)} CẢNH BÁO CẦN KIỂM TRA FILE GỐC ***")
            for w in self.warnings:
                out.append(f"  ⚠ {w}")
            out.append("\n" + "=" * 60 + "\n")
        out.append('\n'.join(self.lines))
        return '\n'.join(out)

    def write_log_file(self, out_dir):
        path = os.path.join(out_dir, f'log_{self.exam_code}.txt')
        with open(path, 'w', encoding='utf-8') as f:
            f.write(self.text())
        return path


# ============================================================
# v4: TÁCH 3 BƯỚC  (load_project -> make_exam_plan -> render_exam)
# ------------------------------------------------------------
# * load_project : đọc & tokenize project gốc MỘT LẦN, ghi nhận đáp án đúng
#                  gốc (marks) của từng câu gốc.
# * make_exam_plan: bốc ngẫu nhiên hoán vị (thứ tự file/cụm/câu + hoán vị
#                  đáp án) cho 1 mã đề sao. Chỉ lưu "kế hoạch", chưa ghi file.
# * render_exam  : ghi file .tex của mã đề theo kế hoạch + marks HIỆN TẠI.
#   => đổi đáp án đúng gốc (marks) chỉ cần render lại, không cần trộn lại,
#      và mọi mã đề sao tự đổi theo. Đây là nền tảng cho web server.
# ============================================================

LETTERS = 'ABCD'


def _identity_perm():
    return [0, 1, 2, 3]


def make_perm(rng, do_shuffle):
    """perm[vị trí mới] = chỉ số đáp án GỐC (0..3) đang đứng ở vị trí đó."""
    p = _identity_perm()
    if do_shuffle:
        rng.shuffle(p)
    return p


def render_choice_q(unit, perm, marks):
    """Dựng lại text câu hỏi với đáp án đã hoán vị; dấu % đặt theo marks
    (tập chỉ số đáp án GỐC đúng) -> hỗ trợ 0, 1 hay nhiều đáp án đúng."""
    pre, macro, args, post = unit['choice']
    new_args = [(args[old][0], old in marks) for old in perm]
    return pre + rebuild_choice(macro, new_args) + post


def new_correct_positions(perm, marks):
    return [p for p, old in enumerate(perm) if old in marks]


def letters_str(idx_list):
    return ''.join(LETTERS[i] for i in sorted(idx_list))


class Project:
    """Kết quả nạp project gốc."""
    pass


def load_project(source_dir):
    struct_path = os.path.join(source_dir, '1.tex')
    if not os.path.isfile(struct_path):
        existing = os.listdir(source_dir) if os.path.isdir(source_dir) else []
        raise FileNotFoundError(
            f"Không tìm thấy '1.tex' trong thư mục:\n  {os.path.abspath(source_dir)}\n"
            f"Danh sách file/thư mục hiện có tại đó:\n  " + ', '.join(existing[:30]) +
            "\n-> Kiểm tra lại đường dẫn bạn nhập (đừng để dấu \" bao quanh, "
            "đừng để dấu \\ cuối chuỗi bị Python hiểu nhầm)."
        )
    P = Project()
    P.source_dir = source_dir
    struct_text = open(struct_path, encoding='utf-8').read()
    P.orig_chunks = build_structure_chunks(parse_structure_file(struct_text))
    P.contents = {}
    P.questions = []          # danh sách câu gốc (theo thứ tự gốc), phần tử = dict
    P.marks = {}              # orig_q -> set(chỉ số đáp án gốc đúng)
    P.warnings = []
    q_counter = 0
    for c in P.orig_chunks:
        if c['path'] is None:
            continue
        raw = open(os.path.join(source_dir, rel_path(c['path'])), encoding='utf-8').read()
        try:
            segments = tokenize_content(raw)
            units, file_ns = build_units(segments)
        except BraceMismatchError as e:
            raise BraceMismatchError(f"[Lỗi khi đọc file: {rel_path(c['path'])}]\n{e}") from None
        for u in units:
            if u['kind'] not in ('q', 'errorq'):
                continue
            q_counter += 1
            u['orig_q_num'] = q_counter
            text = u['seg']['text']
            u['choice'] = None
            marks = set()
            where = f"[{c['path']}.tex] Câu gốc #{q_counter}"
            snip = f"Nội dung: \"{_short_snippet(text)}\""
            if u['kind'] == 'errorq':
                letter = get_errorq_correct_letter(text)
                if letter is None:
                    P.warnings.append(f"{where} (errorq) KHÔNG tìm thấy đáp án đúng "
                                      f"(thiếu ' %A/%B/%C/%D' cuối dòng). {snip}")
                else:
                    marks.add(LETTERS.index(letter))
            else:
                res = parse_choice_in_text(text)
                if res is None:
                    P.warnings.append(f"{where} KHÔNG tìm thấy macro "
                                      f"\\choiceOne/\\choiceTwo/\\choiceFour hợp lệ. {snip}")
                else:
                    u['choice'] = res
                    marks = {i for i, (_, hp) in enumerate(res[2]) if hp}
                    if not marks:
                        P.warnings.append(f"{where} KHÔNG có đáp án nào được đánh dấu đúng "
                                          f"(thiếu '%' ngay sau '}}'). {snip}")
                    elif len(marks) > 1:
                        P.warnings.append(f"{where} có TỚI {len(marks)} đáp án cùng bị đánh dấu "
                                          f"đúng ({letters_str(marks)}). {snip}")
            P.marks[q_counter] = marks
            P.questions.append({'n': q_counter, 'path': c['path'], 'kind': u['kind'],
                                'snippet': _short_snippet(text, 110)})
        P.contents[c['path']] = {'units': units, 'file_ns': file_ns}
    P.n_questions = q_counter
    P.main_text = open(os.path.join(source_dir, 'main.tex'), encoding='utf-8').read()
    P.preamble_text = open(os.path.join(source_dir, 'preamble.tex'), encoding='utf-8').read()
    P.macros_text = open(os.path.join(source_dir, 'macros.tex'), encoding='utf-8').read()
    return P


def make_exam_plan(P, exam_index, rng, total_exams=1):
    """Bốc ngẫu nhiên 1 mã đề sao. Trả về plan (chưa ghi file)."""
    suffix = f"{rng.randint(0, 99):02d}"
    exam_code = f"{exam_index}{suffix}"
    tracker = ProgressTracker(exam_code, exam_index, total_exams)
    print(f"\n[{exam_index}/{total_exams}] Đang tạo mã đề {exam_code}...")
    for w in P.warnings:
        tracker.warn(w)

    key = lambda c: c['rgroup_tag']
    lab = lambda c: c['path'] or '(text)'
    before_map, before_order = snapshot_regions(P.orig_chunks, key, lab)
    tracker.struct_regions(before_map, before_order)
    shuffled_chunks = shuffle_structure_chunks(P.orig_chunks, rng)
    after_map, _ = snapshot_regions(shuffled_chunks, key, lab)
    tracker.struct_regions_after(after_map, before_order)
    print(f"  Đã đảo {len(before_order)} cụm ở cấp file cấu trúc (\\randomgroup)")

    files, items = [], []
    total_files = sum(1 for c in shuffled_chunks if c['path'] is not None)
    file_idx = 0
    for c in shuffled_chunks:
        if c['path'] is None:
            files.append({'path': None, 'prefix': c['prefix'], 'units': [], 'perms': {}})
            continue
        file_idx += 1
        tracker.file_start(file_idx, total_files, c['path'])
        entry = P.contents[c['path']]
        units = list(entry['units'])          # bản sao: KHÔNG đụng vào dữ liệu gốc

        g_before, g_order = snapshot_regions(units, lambda u: u['group_tag'], _label_unit)
        shuffle_within_groups(units, rng)
        g_after, _ = snapshot_regions(units, lambda u: u['group_tag'], _label_unit)
        for tag in g_order:
            tracker.group_region(c['path'], tag, g_before[tag], g_after[tag])

        chunks = make_chunks(units)
        r_before, r_order = snapshot_regions(chunks, lambda ch: ch[0]['rgroup_tag'], _label_chunk)
        units = shuffle_chunks_by_randomgroup(chunks, rng)
        r_after, _ = snapshot_regions(make_chunks(units), lambda ch: ch[0]['rgroup_tag'], _label_chunk)
        for tag in r_order:
            tracker.rgroup_region(c['path'], tag, r_before[tag], r_after[tag])

        perms = {}
        for u in units:
            if u['kind'] in ('q', 'errorq'):
                do_shuffle = (u['kind'] == 'q' and u['choice'] is not None
                              and not (entry['file_ns'] or u['noshuffle']))
                perms[u['orig_q_num']] = make_perm(rng, do_shuffle)
                items.append({'new': len(items) + 1, 'orig': u['orig_q_num'],
                              'perm': perms[u['orig_q_num']]})
        files.append({'path': c['path'], 'prefix': c['prefix'], 'units': units, 'perms': perms})

    return {'index': exam_index, 'code': exam_code, 'files': files, 'items': items,
            'tracker': tracker}


def render_exam(P, plan, output_root):
    """Ghi file .tex + bảng đáp án .txt của 1 mã đề theo marks HIỆN TẠI."""
    out_dir = os.path.join(output_root, plan['code'])
    os.makedirs(out_dir, exist_ok=True)
    struct_out = []
    for f in plan['files']:
        struct_out.append(f['prefix'])
        if f['path'] is None:
            continue
        struct_out.append(r'\inputspace{' + f['path'] + '}')
        rendered = []
        for u in f['units']:
            rendered.append(''.join(s['text'] for s in u['prefix']))
            if u['kind'] == 'q':
                oq = u['orig_q_num']
                if u['choice'] is not None:
                    rendered.append(render_choice_q(u, f['perms'][oq], P.marks[oq]))
                else:
                    rendered.append(u['seg']['text'])
            elif u['kind'] == 'errorq':
                text = u['seg']['text']
                m = P.marks[u['orig_q_num']]
                if len(m) == 1:
                    text = re.sub(r'%([A-D])', '%' + LETTERS[next(iter(m))], text, count=1)
                rendered.append(text)
            elif u['kind'] == 'header':
                rendered.append(u['seg']['text'])
        out_file = os.path.join(out_dir, rel_path(f['path']))
        os.makedirs(os.path.dirname(out_file), exist_ok=True)
        with open(out_file, 'w', encoding='utf-8') as fh:
            fh.write(''.join(rendered))

    idx = plan['index']
    with open(os.path.join(out_dir, f'{idx}.tex'), 'w', encoding='utf-8') as fh:
        fh.write(''.join(struct_out))
    main_text = re.sub(r'\\inputspace\{1\}', r'\\inputspace{%d}' % idx, P.main_text, count=1)
    with open(os.path.join(out_dir, 'main.tex'), 'w', encoding='utf-8') as fh:
        fh.write(main_text)
    pre = re.sub(r'\\newcommand\{\\examcode\}\{[^}]*\}',
                 r'\\newcommand{\\examcode}{%s}' % plan['code'], P.preamble_text)
    with open(os.path.join(out_dir, 'preamble.tex'), 'w', encoding='utf-8') as fh:
        fh.write(pre)
    with open(os.path.join(out_dir, 'macros.tex'), 'w', encoding='utf-8') as fh:
        fh.write(P.macros_text)

    lines = [f"STT, GOC, {plan['code']}", '']
    for it in plan['items']:
        orig = letters_str(P.marks[it['orig']]) or 'None'
        new = letters_str(new_correct_positions(it['perm'], P.marks[it['orig']])) or 'None'
        lines.append(f"{it['new']}, {it['orig']}-{orig}, {it['new']}-{new}")
    with open(os.path.join(out_dir, f"answer_table_{plan['code']}.txt"), 'w', encoding='utf-8') as fh:
        fh.write('\n'.join(lines))
    plan['tracker'].write_log_file(out_dir)
    return out_dir


# ------------------------------------------------------------
# BẢNG ĐÁP ÁN THEO ĐỀ GỐC (CSV mapping + Excel)
# ------------------------------------------------------------

def plan_by_orig(plan):
    """orig_q -> item (new STT, perm)."""
    return {it['orig']: it for it in plan['items']}


def mapping_cell(item, marks):
    """'2-DCA*B': STT ở đề sao - 4 chữ cái GỐC đứng ở vị trí A,B,C,D của đề
    sao; dấu * ngay sau chữ cái là đáp án đúng (có thể nhiều *, hoặc không có)."""
    s = ''.join(LETTERS[old] + ('*' if old in marks else '') for old in item['perm'])
    return f"{item['new']}-{s}"


def build_mapping_csv(P, plans):
    n = P.n_questions
    rows = [['MaDe'] + [f'Cau{i}' for i in range(1, n + 1)]]
    for plan in plans:
        by = plan_by_orig(plan)
        rows.append([plan['code']] + [mapping_cell(by[q], P.marks[q]) for q in range(1, n + 1)])
    return '\r\n'.join(';'.join(r) for r in rows) + '\r\n'


def write_mapping_csv(P, plans, path):
    with open(path, 'w', encoding='utf-8', newline='') as f:
        f.write(build_mapping_csv(P, plans))
    return path


def build_xlsx_workbook(P, plans):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter as CL

    yellow = PatternFill('solid', fgColor='FFFF00')
    head_fill = PatternFill('solid', fgColor='D9E1F2')
    thin, thick = Side(style='thin', color='999999'), Side(style='medium', color='000000')
    ctr = Alignment(horizontal='center', vertical='center')
    bold = Font(bold=True)
    n = P.n_questions
    blocks = [('Gốc', None)] + [(p['code'], p) for p in plans]

    # ---- Sheet 1: đúng như ảnh mẫu (Gốc | mã đề ...; STT A B C D; ô đúng tô vàng) ----
    wb = Workbook()
    ws = wb.active
    ws.title = 'Theo câu gốc'
    for b, (title, plan) in enumerate(blocks):
        c0 = 1 + 5 * b
        ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + 4)
        ws.cell(1, c0, title if plan is None else int(title) if title.isdigit() else title)
        for k, h in enumerate(['STT', 'A', 'B', 'C', 'D']):
            ws.cell(2, c0 + k, h)
        for r in (1, 2):
            for k in range(5):
                cell = ws.cell(r, c0 + k)
                cell.fill, cell.font, cell.alignment = head_fill, bold, ctr
        ws.column_dimensions[CL(c0)].width = 7
        for k in range(1, 5):
            ws.column_dimensions[CL(c0 + k)].width = 5
    by_plan = [plan_by_orig(p) for _, p in blocks[1:]]
    for q in range(1, n + 1):
        r = 2 + q
        ws.cell(r, 1, q)
        for p in P.marks[q]:
            ws.cell(r, 2 + p).fill = yellow
        for b, by in enumerate(by_plan, start=1):
            c0 = 1 + 5 * b
            it = by[q]
            ws.cell(r, c0, it['new'])
            for pos in new_correct_positions(it['perm'], P.marks[q]):
                ws.cell(r, c0 + 1 + pos).fill = yellow
    last_row = 2 + n
    for b in range(len(blocks)):
        c0 = 1 + 5 * b
        for r in range(1, last_row + 1):
            for k in range(5):
                cell = ws.cell(r, c0 + k)
                cell.alignment = ctr
                cell.border = Border(left=thick if k == 0 else thin, right=thin,
                                     top=thin, bottom=thin)
    ws.freeze_panes = 'F3'

    # ---- Sheet 2/3: dạng chữ cái, giống file đáp án cũ ----
    def letters_for(item, q):
        return letters_str(new_correct_positions(item['perm'], P.marks[q]))

    s2 = wb.create_sheet('Theo STT trên đề')
    s2.append(['Câu'] + [p['code'] for p in plans])
    for stt in range(1, n + 1):
        row = [stt]
        for plan in plans:
            it = plan['items'][stt - 1]
            row.append(letters_for(it, it['orig']))
        s2.append(row)
    s3 = wb.create_sheet('Gộp theo câu gốc')
    s3.append(['Câu (gốc)'] + [p['code'] for p in plans])
    for q in range(1, n + 1):
        s3.append([q] + [letters_for(by[q], q) for by in by_plan])
    for w in (s2, s3):
        for cell in w[1]:
            cell.fill, cell.font, cell.alignment = head_fill, bold, ctr
        w.freeze_panes = 'B2'
    for plan in plans:
        s = wb.create_sheet(f"Mã đề {plan['code']}")
        s.append(['STT (câu in trên đề)', 'Đáp án', 'Câu gốc tương ứng'])
        for it in plan['items']:
            s.append([it['new'], letters_for(it, it['orig']), it['orig']])
        for cell in s[1]:
            cell.fill, cell.font, cell.alignment = head_fill, bold, ctr
        s.freeze_panes = 'A2'
        for col, w in zip('ABC', (20, 12, 18)):
            s.column_dimensions[col].width = w
    return wb


def write_xlsx(P, plans, path):
    build_xlsx_workbook(P, plans).save(path)
    return path


def write_all_tables(P, plans, output_root):
    """Ghi mapping_ans_merged.csv + dapan.xlsx vào output_root."""
    csv_path = write_mapping_csv(P, plans, os.path.join(output_root, 'mapping_ans_merged.csv'))
    xlsx_path = write_xlsx(P, plans, os.path.join(output_root, 'dapan.xlsx'))
    return csv_path, xlsx_path


# ============================================================
# CLI
# ============================================================

def _clean_path(p):
    p = p.strip().strip('"').strip("'")
    return p


def main():
    print("=" * 55)
    print("  EXAM MIXER v4 — trộn đề + xuất bảng đáp án theo đề gốc")
    print("=" * 55)

    script_dir = os.path.dirname(os.path.abspath(sys.argv[0])) or '.'
    default_prompt = f"Thư mục chứa đề gốc (Enter = '{script_dir}'): "
    source_dir = _clean_path(input(default_prompt)) or script_dir
    print(f"  -> Dùng thư mục nguồn: {os.path.abspath(source_dir)}")

    n_raw = input("Số lượng mã đề cần sinh (Enter = 1): ").strip()
    n = int(n_raw) if n_raw else 1
    output_root = _clean_path(input("Thư mục output (Enter = 'output'): ")) or f'{script_dir}/output'
    os.makedirs(output_root, exist_ok=True)

    if n > 9:
        print(f"  ⚠ Lưu ý: bạn yêu cầu {n} mã đề (STT > 9). Mã đề sẽ có dạng "
              f"'<STT><2 số random>' nên từ STT 10 trở đi mã đề sẽ có 4 chữ số.")

    rng = random.SystemRandom()
    try:
        P = load_project(source_dir)
        print(f"  Đã nạp {P.n_questions} câu gốc, {len(P.warnings)} cảnh báo.")
        plans = []
        for k in range(1, n + 1):
            plan = make_exam_plan(P, k, rng, total_exams=n)
            out_dir = render_exam(P, plan, output_root)
            plan['tracker'].finish(len(plan['items']), out_dir)
            plans.append(plan)
    except BraceMismatchError as e:
        print(f"\n❌ DỪNG LẠI — lỗi cấu trúc ngoặc {{}} trong file nội dung:\n{e}\n")
        sys.exit(1)
    except FileNotFoundError as e:
        print(f"\n❌ DỪNG LẠI — không tìm thấy file:\n{e}\n")
        sys.exit(1)

    csv_path, xlsx_path = write_all_tables(P, plans, output_root)
    print("\nHoàn tất! Các mã đề đã sinh:", ', '.join(p['code'] for p in plans))
    print(f"  Mapping đáp án : {csv_path}")
    print(f"  Bảng đáp án    : {xlsx_path}")
    print("=" * 55)


if __name__ == '__main__':
    main()
